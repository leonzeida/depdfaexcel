#!/usr/bin/env python3
"""Acceso a la base de datos de Zeid Medical: precios de referencia del
comparador de precios de proveedores, y el historial de licitaciones
(oferta -> resultado real -> costo real) que vincula el comparador (F.41)
con las Actas de Preadjudicación (F.43).

Se mantiene separado de f41_a_excel.py a propósito: ese módulo no debería
necesitar una base de datos para poder usarse por línea de comandos.

Requiere una variable de entorno:
- DATABASE_URL: connection string de Postgres (ej. la que da Neon), con
  el formato postgres://usuario:contraseña@host/db?sslmode=require.
"""

import os
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from f41_a_excel import FMT_MONEDA_ARS, clave_item, normalizar_descripcion, parsear_fecha_apertura

ZONA_HORARIA_ARGENTINA = ZoneInfo("America/Argentina/Buenos_Aires")

_CREAR_TABLA = """
CREATE TABLE IF NOT EXISTS precios_referencia (
    codigo TEXT NOT NULL,
    descripcion TEXT NOT NULL,
    ultimo_precio NUMERIC(12,2),
    actualizado DATE,
    PRIMARY KEY (codigo, descripcion)
);
"""

# ADD COLUMN IF NOT EXISTS para no romper la tabla ya creada en producción
# (antes de agregar "porcentaje", el CREATE TABLE de arriba ya se había
# corrido ahí, así que un CREATE TABLE nuevo no le agrega la columna sola).
_AGREGAR_COLUMNA_PORCENTAJE = """
ALTER TABLE precios_referencia ADD COLUMN IF NOT EXISTS porcentaje NUMERIC(12,2);
"""

# Mismo criterio: se agrega después de que la tabla ya existía en
# producción sin esta columna.
_AGREGAR_COLUMNA_MEJOR_PROVEEDOR = """
ALTER TABLE precios_referencia ADD COLUMN IF NOT EXISTS mejor_proveedor TEXT;
"""

_UPSERT = """
INSERT INTO precios_referencia (codigo, descripcion, ultimo_precio, porcentaje, mejor_proveedor, actualizado)
VALUES (%s, %s, %s, %s, %s, %s)
ON CONFLICT (codigo, descripcion) DO UPDATE SET
    -- "actualizado" ya no es "cuándo se tocó el botón Guardar" -- es la
    -- fecha de APERTURA real del documento (F.41 o F.43) de donde salió
    -- ese precio (ver guardar_precios_referencia). Por eso el pisado ya
    -- no es incondicional: si lo que ya había guardado es de una
    -- licitación más reciente que la que se acaba de cargar (ej. el
    -- padre carga un F.43 atrasado después de uno más nuevo), se
    -- mantiene el dato más nuevo en vez de taparlo con uno viejo. Los 4
    -- campos se mueven juntos (mismo CASE) porque son "un mismo precio
    -- de un mismo documento" -- antes mezclaban criterios distintos
    -- (ultimo_precio/actualizado se pisaban siempre, porcentaje/
    -- mejor_proveedor sólo con COALESCE).
    ultimo_precio = CASE
        WHEN precios_referencia.actualizado IS NULL OR EXCLUDED.actualizado >= precios_referencia.actualizado
        THEN EXCLUDED.ultimo_precio ELSE precios_referencia.ultimo_precio END,
    porcentaje = CASE
        WHEN precios_referencia.actualizado IS NULL OR EXCLUDED.actualizado >= precios_referencia.actualizado
        THEN COALESCE(EXCLUDED.porcentaje, precios_referencia.porcentaje) ELSE precios_referencia.porcentaje END,
    mejor_proveedor = CASE
        WHEN precios_referencia.actualizado IS NULL OR EXCLUDED.actualizado >= precios_referencia.actualizado
        THEN COALESCE(EXCLUDED.mejor_proveedor, precios_referencia.mejor_proveedor) ELSE precios_referencia.mejor_proveedor END,
    actualizado = CASE
        WHEN precios_referencia.actualizado IS NULL OR EXCLUDED.actualizado >= precios_referencia.actualizado
        THEN EXCLUDED.actualizado ELSE precios_referencia.actualizado END;
"""

# Una fila por (item, licitacion), que se va completando en 3 etapas a
# medida que pasa el tiempo en vez de pisarse (a diferencia de
# precios_referencia, que es "la ultima foto" de cada item):
#   1. oferta (comparador F.41): que precio calculo Zeid, con que proveedor propio.
#   2. resultado (F.43): a quien se lo adjudicaron en la realidad y a que precio.
#   3. costo real (F.43, items ganados): cuanto salio comprarlo de verdad.
_CREAR_TABLA_HISTORIAL = """
CREATE TABLE IF NOT EXISTS historial_licitaciones (
    codigo TEXT NOT NULL,
    descripcion_normalizada TEXT NOT NULL,
    expediente TEXT NOT NULL,
    descripcion TEXT,
    contratacion TEXT,
    cantidad NUMERIC(12,2),
    precio_ofertado NUMERIC(12,2),
    porcentaje_ofertado NUMERIC(12,2),
    proveedor_elegido TEXT,
    fecha_oferta DATE,
    gano_zeid BOOLEAN,
    precio_adjudicado NUMERIC(12,2),
    proveedor_ganador TEXT,
    fecha_adjudicacion DATE,
    costo_real NUMERIC(12,2),
    precio_venta_real NUMERIC(12,2),
    fecha_costo DATE,
    PRIMARY KEY (codigo, descripcion_normalizada, expediente)
);
"""

# Agregada después de que la tabla ya existía en producción, por eso
# ALTER TABLE en vez de meterla directo en el CREATE TABLE de arriba.
# "Renglón" es el número de línea oficial del ítem DENTRO de un
# expediente puntual -- a diferencia de descripcion_normalizada (que es
# una aproximación de texto), es el identificador que usa el propio
# trámite y es estable entre el F.41 y el F.43 de una misma licitación
# (confirmado comparando renglón por renglón en 3 expedientes reales:
# coincide 100% de las veces, excepto los renglones declarados
# "desiertos" en el F.43, que no tienen código). Se usa para unir mejor
# la etapa 1 (oferta) con la etapa 2 (resultado) cuando las descripciones
# del F.41 y el F.43 no son textualmente idénticas -- ver
# _resolver_clave_existente. No reemplaza la PRIMARY KEY existente (eso
# requeriría migrar datos reales ya guardados que no tienen renglón).
_AGREGAR_COLUMNA_RENGLON = """
ALTER TABLE historial_licitaciones ADD COLUMN IF NOT EXISTS renglon INTEGER;
"""

# COALESCE(EXCLUDED.x, tabla.x) en vez de pisar directo: las etapas 1/2/3
# se guardan en momentos distintos (a veces días de diferencia) y cada
# upsert solo trae sus propias columnas - sin esto, guardar la etapa 2
# pisaría con NULL lo que ya se había guardado en la etapa 1, y viceversa.
_UPSERT_OFERTA = """
INSERT INTO historial_licitaciones (
    codigo, descripcion_normalizada, expediente, renglon, descripcion, contratacion, cantidad,
    precio_ofertado, porcentaje_ofertado, proveedor_elegido, fecha_oferta
) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
ON CONFLICT (codigo, descripcion_normalizada, expediente) DO UPDATE SET
    renglon = COALESCE(EXCLUDED.renglon, historial_licitaciones.renglon),
    descripcion = COALESCE(EXCLUDED.descripcion, historial_licitaciones.descripcion),
    contratacion = COALESCE(EXCLUDED.contratacion, historial_licitaciones.contratacion),
    cantidad = COALESCE(EXCLUDED.cantidad, historial_licitaciones.cantidad),
    precio_ofertado = COALESCE(EXCLUDED.precio_ofertado, historial_licitaciones.precio_ofertado),
    porcentaje_ofertado = COALESCE(EXCLUDED.porcentaje_ofertado, historial_licitaciones.porcentaje_ofertado),
    proveedor_elegido = COALESCE(EXCLUDED.proveedor_elegido, historial_licitaciones.proveedor_elegido),
    fecha_oferta = COALESCE(EXCLUDED.fecha_oferta, historial_licitaciones.fecha_oferta);
"""

_UPSERT_ADJUDICACION = """
INSERT INTO historial_licitaciones (
    codigo, descripcion_normalizada, expediente, renglon, descripcion, contratacion, cantidad,
    gano_zeid, precio_adjudicado, proveedor_ganador, fecha_adjudicacion,
    costo_real, precio_venta_real, fecha_costo
) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
ON CONFLICT (codigo, descripcion_normalizada, expediente) DO UPDATE SET
    renglon = COALESCE(EXCLUDED.renglon, historial_licitaciones.renglon),
    descripcion = COALESCE(EXCLUDED.descripcion, historial_licitaciones.descripcion),
    contratacion = COALESCE(EXCLUDED.contratacion, historial_licitaciones.contratacion),
    cantidad = COALESCE(EXCLUDED.cantidad, historial_licitaciones.cantidad),
    gano_zeid = COALESCE(EXCLUDED.gano_zeid, historial_licitaciones.gano_zeid),
    precio_adjudicado = COALESCE(EXCLUDED.precio_adjudicado, historial_licitaciones.precio_adjudicado),
    proveedor_ganador = COALESCE(EXCLUDED.proveedor_ganador, historial_licitaciones.proveedor_ganador),
    fecha_adjudicacion = COALESCE(EXCLUDED.fecha_adjudicacion, historial_licitaciones.fecha_adjudicacion),
    costo_real = COALESCE(EXCLUDED.costo_real, historial_licitaciones.costo_real),
    precio_venta_real = COALESCE(EXCLUDED.precio_venta_real, historial_licitaciones.precio_venta_real),
    fecha_costo = COALESCE(EXCLUDED.fecha_costo, historial_licitaciones.fecha_costo);
"""


class ErrorPreciosReferencia(Exception):
    """Cualquier problema al leer o guardar en la base de datos."""


def _conectar():
    url = os.environ.get("DATABASE_URL")
    if not url:
        raise ErrorPreciosReferencia("Falta configurar DATABASE_URL.")
    try:
        conn = psycopg.connect(url)
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo conectar a la base de datos: {type(exc).__name__}: {exc!r}"
        ) from exc
    cur = conn.cursor()
    cur.execute(_CREAR_TABLA)
    cur.execute(_AGREGAR_COLUMNA_PORCENTAJE)
    cur.execute(_AGREGAR_COLUMNA_MEJOR_PROVEEDOR)
    cur.execute(_CREAR_TABLA_HISTORIAL)
    cur.execute(_AGREGAR_COLUMNA_RENGLON)
    cur.close()
    conn.commit()
    return conn


def _resolver_clave_existente(cur, expediente, renglon, codigo, descripcion_normalizada):
    """Si ya hay una fila guardada para este renglón de este expediente
    (por ejemplo la oferta del F.41, guardada antes de cargar el F.43, o
    viceversa), reusa SU (codigo, descripcion_normalizada) para el
    upsert, en vez del que se acaba de calcular -- así las dos etapas se
    unen en la misma fila incluso si el F.41 y el F.43 formatearon la
    descripción de ese ítem de forma distinta. Sin esto, cada documento
    podría terminar creando su propia fila separada para el mismo ítem
    real."""
    if renglon is None:
        return codigo, descripcion_normalizada
    cur.execute(
        "SELECT codigo, descripcion_normalizada FROM historial_licitaciones "
        "WHERE expediente = %s AND renglon = %s LIMIT 1;",
        (expediente, renglon),
    )
    fila = cur.fetchone()
    return (fila[0], fila[1]) if fila else (codigo, descripcion_normalizada)


def leer_precios_referencia() -> dict:
    """Devuelve {(codigo, descripcion_normalizada): {"ultimo_precio": float|None,
    "porcentaje": float|None, "actualizado": str|None (fecha ISO),
    "mejor_proveedor": str|None}}.

    La tabla todavía guarda la descripción tal cual vino (texto crudo,
    sin normalizar) -- distintos documentos pueden haber guardado "el
    mismo" ítem real con un texto apenas distinto (mayúsculas, notas
    "Obs.:", un "X 30" de más), así que puede haber más de una fila
    cruda para la misma clave normalizada (clave_item ya normaliza, ver
    f41_a_excel.py). Se resuelve acá, en memoria, en vez de migrar la
    tabla: se leen todas ordenadas por fecha ascendente y se arma el
    diccionario en ese orden, así si dos filas colisionan en la misma
    clave normalizada, la que se guarda al final (la de fecha más
    reciente) es la que queda."""
    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        cur.execute(
            "SELECT codigo, descripcion, ultimo_precio, porcentaje, actualizado, mejor_proveedor "
            "FROM precios_referencia ORDER BY actualizado ASC NULLS FIRST;"
        )
        filas = cur.fetchall()
        cur.close()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo leer la tabla de precios de referencia: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()

    precios = {}
    for codigo, descripcion, ultimo_precio, porcentaje, actualizado, mejor_proveedor in filas:
        precios[clave_item(codigo, descripcion)] = {
            "ultimo_precio": float(ultimo_precio) if ultimo_precio is not None else None,
            "porcentaje": float(porcentaje) if porcentaje is not None else None,
            "actualizado": actualizado.isoformat() if actualizado is not None else None,
            "mejor_proveedor": mejor_proveedor,
        }
    return precios


def guardar_precios_referencia(items: list, fecha_documento=None):
    """Por cada {"codigo", "descripcion", "ultimo_precio", "porcentaje",
    "mejor_proveedor"}, hace un upsert en la tabla (crea la fila si no
    existía, o actualiza los datos si ya existía). `fecha_documento`: la
    fecha de APERTURA real del documento (F.41 o F.43) del que salen
    estos items -- un solo valor para todo el batch, no por item, porque
    todos vienen del mismo documento. El UPSERT (_UPSERT) sólo pisa un
    precio ya guardado si esta fecha es igual o más nueva que la que ya
    había, así un documento cargado tarde no tapa uno más reciente. Si
    no se pasa nada (o no se pudo parsear la fecha del documento), cae a
    "hoy" como antes."""
    if not items:
        return

    fecha = fecha_documento or datetime.now(ZONA_HORARIA_ARGENTINA).date()
    filas = [
        (
            item["codigo"],
            item["descripcion"],
            item["ultimo_precio"],
            item.get("porcentaje"),
            item.get("mejor_proveedor"),
            fecha,
        )
        for item in items
    ]

    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        cur.executemany(_UPSERT, filas)
        cur.close()
        conn.commit()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo guardar en la tabla de precios de referencia: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()


def guardar_oferta_historial(items: list, encabezado: dict):
    """Guarda la etapa 1 (lo que Zeid ofertó) del historial de
    licitaciones. `items`: lista de {"codigo", "descripcion", "cantidad",
    "precio_ofertado", "porcentaje_ofertado", "proveedor_elegido",
    "renglon"?}. `encabezado`: {"expediente", "contratacion", "apertura"}
    del F.41 cargado. No hace nada si no hay expediente (no se puede
    vincular a una licitación)."""
    expediente = (encabezado or {}).get("expediente") or ""
    if not items or not expediente:
        return

    contratacion = (encabezado or {}).get("contratacion")
    # "fecha_oferta" es la fecha de APERTURA real del F.41, no el día en
    # que se tocó "Guardar" -- mismo criterio que precios_referencia (ver
    # guardar_precios_referencia). Si no se pudo extraer/parsear, cae a
    # hoy como antes.
    fecha_oferta = parsear_fecha_apertura(encabezado.get("apertura")) or datetime.now(ZONA_HORARIA_ARGENTINA).date()

    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        for item in items:
            renglon = item.get("renglon")
            codigo, descripcion_normalizada = _resolver_clave_existente(
                cur, expediente, renglon, item["codigo"], normalizar_descripcion(item["descripcion"])
            )
            cur.execute(
                _UPSERT_OFERTA,
                (
                    codigo,
                    descripcion_normalizada,
                    expediente,
                    renglon,
                    item["descripcion"],
                    contratacion,
                    item.get("cantidad"),
                    item.get("precio_ofertado"),
                    item.get("porcentaje_ofertado"),
                    item.get("proveedor_elegido"),
                    fecha_oferta,
                ),
            )
        cur.close()
        conn.commit()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo guardar el historial de la oferta: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()


def guardar_resultado_adjudicacion(items: list, encabezado: dict):
    """Guarda la etapa 2 (resultado real) y, si vienen, la etapa 3
    (costo real) del historial de licitaciones. `items`: lista de
    {"codigo", "descripcion", "cantidad", "gano_zeid",
    "precio_adjudicado", "proveedor_ganador", "costo_real"?,
    "precio_venta_real"?, "renglon"?}. `encabezado`: {"expediente",
    "contratacion", "apertura"} del F.43 cargado."""
    expediente = (encabezado or {}).get("expediente") or ""
    if not items or not expediente:
        return

    contratacion = (encabezado or {}).get("contratacion")
    hoy = datetime.now(ZONA_HORARIA_ARGENTINA).date()
    # "fecha_adjudicacion" es la fecha de APERTURA real del F.43, no el
    # día en que se tocó "Guardar resultados" (mismo criterio que
    # fecha_oferta arriba). "fecha_costo" sí se deja en "hoy": no es la
    # fecha de ningún documento, es el día en que el padre cargó el costo
    # real a mano.
    fecha_adjudicacion = parsear_fecha_apertura(encabezado.get("apertura")) or hoy

    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        for item in items:
            renglon = item.get("renglon")
            codigo, descripcion_normalizada = _resolver_clave_existente(
                cur, expediente, renglon, item["codigo"], normalizar_descripcion(item["descripcion"])
            )
            tiene_costo = item.get("costo_real") is not None or item.get("precio_venta_real") is not None
            cur.execute(
                _UPSERT_ADJUDICACION,
                (
                    codigo,
                    descripcion_normalizada,
                    expediente,
                    renglon,
                    item["descripcion"],
                    contratacion,
                    item.get("cantidad"),
                    item.get("gano_zeid"),
                    item.get("precio_adjudicado"),
                    item.get("proveedor_ganador"),
                    fecha_adjudicacion,
                    item.get("costo_real"),
                    item.get("precio_venta_real"),
                    hoy if tiene_costo else None,
                ),
            )
        cur.close()
        conn.commit()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo guardar el resultado de la adjudicación: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()


def leer_resumen_proveedores() -> dict:
    """Devuelve {"proveedores": [{"proveedor", "veces_elegido",
    "veces_ganadas", "porcentaje_exito"}, ...] (ordenado por veces
    elegido, descendente), "totales": {"licitaciones_ganadas",
    "costo_total", "venta_total", "margen_total"}}."""
    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        cur.execute(
            "SELECT proveedor_elegido, COUNT(*), COUNT(*) FILTER (WHERE gano_zeid) "
            "FROM historial_licitaciones "
            "WHERE proveedor_elegido IS NOT NULL AND proveedor_elegido <> '' "
            "GROUP BY proveedor_elegido "
            "ORDER BY COUNT(*) DESC;"
        )
        filas_proveedores = cur.fetchall()
        cur.execute(
            "SELECT COUNT(*) FILTER (WHERE gano_zeid), "
            "SUM(costo_real * cantidad) FILTER (WHERE gano_zeid), "
            "SUM(precio_venta_real * cantidad) FILTER (WHERE gano_zeid) "
            "FROM historial_licitaciones;"
        )
        licitaciones_ganadas, costo_total, venta_total = cur.fetchone()
        cur.close()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo leer el resumen de proveedores: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()

    proveedores = [
        {
            "proveedor": proveedor,
            "veces_elegido": veces_elegido,
            "veces_ganadas": veces_ganadas,
            "porcentaje_exito": round(100 * veces_ganadas / veces_elegido, 1) if veces_elegido else 0,
        }
        for proveedor, veces_elegido, veces_ganadas in filas_proveedores
    ]

    costo_total = float(costo_total) if costo_total is not None else 0.0
    venta_total = float(venta_total) if venta_total is not None else 0.0
    return {
        "proveedores": proveedores,
        "totales": {
            "licitaciones_ganadas": licitaciones_ganadas or 0,
            "costo_total": costo_total,
            "venta_total": venta_total,
            "margen_total": venta_total - costo_total,
        },
    }


def leer_historial_por_expediente(expediente: str) -> list:
    """Todas las filas de historial_licitaciones para un expediente
    puntual (sin la columna interna descripcion_normalizada), ordenadas
    por renglón (el orden real del expediente -- cae a código para las
    filas viejas que se guardaron antes de existir esta columna, donde
    renglon es NULL)."""
    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        cur.execute(
            "SELECT codigo, descripcion, cantidad, precio_ofertado, proveedor_elegido, "
            "gano_zeid, precio_adjudicado, proveedor_ganador, costo_real, precio_venta_real, renglon "
            "FROM historial_licitaciones WHERE expediente = %s ORDER BY renglon NULLS LAST, codigo;",
            (expediente,),
        )
        filas = cur.fetchall()
        cur.close()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo leer el historial del expediente: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()

    return [
        {
            "codigo": codigo,
            "descripcion": descripcion,
            "cantidad": float(cantidad) if cantidad is not None else None,
            "precio_ofertado": float(precio_ofertado) if precio_ofertado is not None else None,
            "proveedor_elegido": proveedor_elegido,
            "gano_zeid": gano_zeid,
            "precio_adjudicado": float(precio_adjudicado) if precio_adjudicado is not None else None,
            "proveedor_ganador": proveedor_ganador,
            "costo_real": float(costo_real) if costo_real is not None else None,
            "precio_venta_real": float(precio_venta_real) if precio_venta_real is not None else None,
            "renglon": renglon,
        }
        for (
            codigo, descripcion, cantidad, precio_ofertado, proveedor_elegido,
            gano_zeid, precio_adjudicado, proveedor_ganador, costo_real, precio_venta_real, renglon,
        ) in filas
    ]


def leer_historial_completo() -> list:
    """Todas las filas de historial_licitaciones, todas las columnas
    visibles (sin la interna descripcion_normalizada), para exportar a
    Excel. Devuelve tuplas crudas (no dicts) en el mismo orden que
    ENCABEZADOS_HISTORIAL, listas para volcar directo a una hoja.

    Ordenado por expediente, pero los expedientes en sí van del más
    reciente al más viejo (según fecha_adjudicacion si existe, si no
    fecha_oferta -- ya NO es "cuándo se cargó", es la fecha de apertura
    real del documento, a propósito). Adentro de cada expediente se
    mantiene el orden por renglón de siempre."""
    conn = None
    try:
        conn = _conectar()
        cur = conn.cursor()
        cur.execute(
            "SELECT codigo, renglon, descripcion, expediente, contratacion, cantidad, "
            "precio_ofertado, porcentaje_ofertado, proveedor_elegido, fecha_oferta, "
            "gano_zeid, precio_adjudicado, proveedor_ganador, fecha_adjudicacion, "
            "costo_real, precio_venta_real, fecha_costo "
            "FROM historial_licitaciones "
            "ORDER BY MAX(COALESCE(fecha_adjudicacion, fecha_oferta)) OVER (PARTITION BY expediente) "
            "DESC NULLS LAST, expediente, renglon NULLS LAST, codigo;"
        )
        filas = cur.fetchall()
        cur.close()
    except ErrorPreciosReferencia:
        raise
    except Exception as exc:
        raise ErrorPreciosReferencia(
            f"No se pudo leer el historial completo: {type(exc).__name__}: {exc!r}"
        ) from exc
    finally:
        if conn is not None:
            conn.close()
    return filas


ENCABEZADOS_HISTORIAL = [
    "Codigo", "Renglon", "Descripcion", "Expediente", "Contratacion", "Cantidad",
    "Precio ofertado", "% ofertado", "Proveedor elegido", "Fecha oferta",
    "Gano Zeid", "Precio adjudicado", "Proveedor ganador", "Fecha adjudicacion",
    "Costo real", "Precio venta real", "Fecha costo",
]


def escribir_historial_excel(filas: list, salida: Path):
    """Vuelca a un .xlsx todo el detalle crudo de historial_licitaciones
    (una fila por ítem+licitación), para archivar o analizar aparte.
    `filas` es la lista de tuplas que devuelve leer_historial_completo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Historial"

    fuente_encabezado = Font(bold=True, color="000000")
    relleno_encabezado = PatternFill("solid", fgColor="D4EA6B")
    centrado = Alignment(horizontal="center")

    for col, titulo in enumerate(ENCABEZADOS_HISTORIAL, start=1):
        c = ws.cell(row=1, column=col, value=titulo)
        c.font = fuente_encabezado
        c.fill = relleno_encabezado
        c.alignment = centrado

    columnas_moneda = (7, 12, 15, 16)  # Precio ofertado, Precio adjudicado, Costo real, Precio venta real
    columnas_fecha = (10, 14, 17)
    columna_booleana = 11  # Gano Zeid
    fila_excel = 2
    for fila in filas:
        for col, valor in enumerate(fila, start=1):
            if col in columnas_fecha and valor is not None:
                valor = valor.isoformat()
            elif col == columna_booleana and valor is not None:
                valor = "Si" if valor else "No"
            c = ws.cell(row=fila_excel, column=col, value=valor)
            c.alignment = Alignment(horizontal="left", wrap_text=True) if col == 3 else centrado
            if col in columnas_moneda and valor is not None:
                c.number_format = FMT_MONEDA_ARS
        fila_excel += 1

    anchos = {
        "A": 14, "B": 10, "C": 46, "D": 14, "E": 14, "F": 10, "G": 14,
        "H": 10, "I": 24, "J": 12, "K": 10, "L": 14, "M": 24, "N": 14, "O": 12, "P": 14, "Q": 12,
    }
    for letra, ancho in anchos.items():
        ws.column_dimensions[letra].width = ancho

    wb.save(salida)
